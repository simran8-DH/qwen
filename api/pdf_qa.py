"""
File / PDF Q&A (NO RAG — no FAISS, no embeddings).

APIs:
  POST /api/v1/file-to-base64   — upload file → { filename, file_base64 }
  POST /api/v1/file-qa-base64   — JSON { filename, file_base64, question } → answer
  POST /api/v1/file-qa-form     — form-data file + question → answer (kept)

Flow:
  1) Client encodes file as base64 (or call /file-to-base64)
  2) POST /file-qa-base64 with filename + file_base64 + question
  3) Server extracts text → stuffs into prompt → Qwen answers

Technique: prompt stuffing. If doc > N_CTX → truncated=true.
"""

from __future__ import annotations

import asyncio
import base64
import binascii
import csv
import io
import json
import uuid
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader

from config import settings
from llm.engine import generate, get_llm

router = APIRouter(prefix="/api/v1", tags=["File Q&A"])

_CHARS_PER_TOKEN = 3.5
_PROMPT_OVERHEAD_TOKENS = 400
_MAX_UPLOAD_BYTES = 20 * 1024 * 1024

_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl",
    ".log", ".xml", ".html", ".htm", ".yaml", ".yml", ".ini", ".cfg",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs", ".go", ".rs",
    ".sql", ".rtf",
}


class FileQaResponse(BaseModel):
    id: str
    filename: str
    file_type: str
    pages: int | None = None
    chars_used: int   # document text actually sent to Qwen
    truncated: bool   # true if document was cut to fit N_CTX
    question: str
    answer: str
    model: str = "qwen-local"


class FileToBase64Response(BaseModel):
    """POST /file-to-base64 response."""
    filename: str
    file_base64: str
    size_bytes: int


class FileQaBase64Request(BaseModel):
    """POST /file-qa-base64 body."""
    filename: str = Field(..., description="e.g. document.pdf")
    file_base64: str = Field(..., description="Base64-encoded file bytes")
    question: str
    max_tokens: int | None = Field(None, ge=1, le=2048)


def _ensure_model() -> None:
    try:
        get_llm()
    except ModuleNotFoundError as exc:
        if "llama_cpp" in str(exc):
            raise HTTPException(
                503,
                "llama-cpp-python not installed. Start with: .venv\\Scripts\\python.exe main.py",
            ) from exc
        raise
    except (RuntimeError, OSError, FileNotFoundError) as exc:
        raise HTTPException(503, f"LLM failed to load: {exc}") from exc


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> tuple[str, int | None]:
    """Text-layer PDF only — scanned/image PDFs need OCR (not supported)."""
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(400, f"Invalid or unreadable PDF: {exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise HTTPException(400, f"PDF is encrypted/password-protected: {exc}") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(part.strip() for part in pages if part and part.strip()).strip()
    return text, len(reader.pages)


def _extract_docx(data: bytes) -> tuple[str, int | None]:
    try:
        from docx import Document
    except ModuleNotFoundError as exc:
        raise HTTPException(503, "python-docx not installed") from exc
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(400, f"Invalid or unreadable DOCX: {exc}") from exc

    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip(), None


def _extract_xlsx(data: bytes) -> tuple[str, int | None]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise HTTPException(503, "openpyxl not installed") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Invalid or unreadable Excel file: {exc}") from exc

    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts).strip(), None


def _extract_csv_like(data: bytes, *, delimiter: str = ",") -> tuple[str, int | None]:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [" | ".join(cell.strip() for cell in row if cell and cell.strip()) for row in reader]
    return "\n".join(r for r in rows if r).strip(), None


def _extract_json(data: bytes) -> tuple[str, int | None]:
    text = _decode_text(data).strip()
    try:
        return json.dumps(json.loads(text), ensure_ascii=False, indent=2), None
    except json.JSONDecodeError:
        return text, None


def _extract_plain_text(data: bytes) -> tuple[str, int | None]:
    return _decode_text(data).strip(), None


def _detect_kind(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if data.startswith(b"%PDF") or ext == ".pdf":
        return "pdf"
    if ext == ".docx" or (data[:2] == b"PK" and ext == ".docx"):
        return "docx"
    if ext in {".xlsx", ".xlsm"}:
        return "xlsx"
    if ext == ".csv":
        return "csv"
    if ext == ".tsv":
        return "tsv"
    if ext in {".json", ".jsonl"}:
        return "json"
    if ext in _TEXT_EXTENSIONS:
        return "text"
    if data.startswith(b"{") or data.startswith(b"["):
        return "json"
    if b"\x00" not in data[:4096]:
        return "text"
    raise HTTPException(400, f"Unsupported file type '{ext or 'unknown'}'")


def extract_file_text(filename: str, data: bytes) -> tuple[str, str, int | None]:
    kind = _detect_kind(filename, data)
    if kind == "pdf":
        text, pages = _extract_pdf(data)
    elif kind == "docx":
        text, pages = _extract_docx(data)
    elif kind == "xlsx":
        text, pages = _extract_xlsx(data)
    elif kind == "csv":
        text, pages = _extract_csv_like(data, delimiter=",")
    elif kind == "tsv":
        text, pages = _extract_csv_like(data, delimiter="\t")
    elif kind == "json":
        text, pages = _extract_json(data)
    else:
        text, pages = _extract_plain_text(data)
    return text, kind, pages


def _fit_to_context(text: str, question: str) -> tuple[str, bool]:
    """truncated=false → full text used; true → beginning only (doc too long for N_CTX)."""
    budget_tokens = max(
        512,
        settings.n_ctx - settings.max_tokens - _PROMPT_OVERHEAD_TOKENS - max(64, len(question) // 3),
    )
    max_chars = int(budget_tokens * _CHARS_PER_TOKEN)
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


def _build_user_prompt(doc_text: str, question: str, *, truncated: bool, filename: str) -> str:
    note = (
        "\n\n[Note: the file was truncated to fit the model context. Answer only from the text below.]"
        if truncated
        else ""
    )
    return (
        "Answer the user's question using ONLY the file content below. "
        "If the answer is not in the file, say you cannot find it in the document. "
        "Be concise and factual.\n\n"
        f"Filename: {filename}\n"
        f"--- FILE START ---{note}\n"
        f"{doc_text}\n"
        f"--- FILE END ---\n\n"
        f"Question: {question}"
    )


def _decode_base64_file(file_base64: str) -> bytes:
    raw = (file_base64 or "").strip()
    if not raw:
        raise HTTPException(400, "file_base64 is required")
    if "," in raw and raw.lower().startswith("data:"):
        raw = raw.split(",", 1)[1]
    raw = "".join(raw.split())
    try:
        return base64.b64decode(raw, validate=False)
    except (binascii.Error, ValueError) as exc:
        raise HTTPException(400, f"Invalid base64: {exc}") from exc


async def _answer_file(
    *,
    filename: str,
    data: bytes,
    question: str,
    max_tokens: int | None,
) -> FileQaResponse:
    """Shared pipeline: extract text → fit context → call Qwen."""
    if not question.strip():
        raise HTTPException(400, "question is required")
    if not data:
        raise HTTPException(400, "Uploaded file is empty")
    if len(data) > _MAX_UPLOAD_BYTES:
        raise HTTPException(400, "File too large (max 20 MB)")

    filename = (filename or "upload.bin").strip() or "upload.bin"
    text, file_type, pages = await asyncio.to_thread(extract_file_text, filename, data)
    if not text:
        raise HTTPException(
            400,
            "No extractable text in this file. Scanned/image PDFs need OCR (not supported).",
        )

    doc_text, truncated = _fit_to_context(text, question.strip())
    await asyncio.to_thread(_ensure_model)

    user_prompt = _build_user_prompt(
        doc_text, question.strip(), truncated=truncated, filename=filename
    )
    try:
        answer = await asyncio.to_thread(
            lambda: generate(
                user=user_prompt,
                system="You are a careful document assistant. Use only the provided file text.",
                max_tokens=max_tokens or settings.max_tokens,
                temperature=0.2,
            )
        )
    except Exception as exc:
        raise HTTPException(500, f"File Q&A failed: {exc}") from exc

    return FileQaResponse(
        id=f"fileqa_{uuid.uuid4().hex[:12]}",
        filename=filename,
        file_type=file_type,
        pages=pages,
        chars_used=len(doc_text),
        truncated=truncated,
        question=question.strip(),
        answer=answer,
        model="qwen-local",
    )


@router.post("/file-to-base64", response_model=FileToBase64Response)
async def file_to_base64(file: UploadFile = File(...)) -> FileToBase64Response:
    """
    Step 1 — convert uploaded file to base64.
    Returns { filename, file_base64, size_bytes }.
    """
    filename = (file.filename or "upload.bin").strip() or "upload.bin"
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty")
    if len(data) > _MAX_UPLOAD_BYTES:
        raise HTTPException(400, "File too large (max 20 MB)")
    return FileToBase64Response(
        filename=filename,
        file_base64=base64.b64encode(data).decode("ascii"),
        size_bytes=len(data),
    )


@router.post("/file-qa-base64", response_model=FileQaResponse)
async def file_qa_base64(req: FileQaBase64Request) -> FileQaResponse:
    """
    Step 2 — ask Qwen using base64 file content (JSON, no multipart file).

    Body:
      {
        "filename": "doc.pdf",
        "file_base64": "...",
        "question": "...",
        "max_tokens": 256
      }
    """
    data = _decode_base64_file(req.file_base64)
    return await _answer_file(
        filename=req.filename,
        data=data,
        question=req.question,
        max_tokens=req.max_tokens,
    )


@router.post("/file-qa-form", response_model=FileQaResponse)
async def file_qa_form(
    file: UploadFile = File(...),
    question: str = Form(...),
    max_tokens: int | None = Form(None, ge=1, le=2048),
) -> FileQaResponse:
    """
    One-step form-data: file + question together.
    Kept as-is; LAN firewalls may block multipart uploads.
    """
    filename = (file.filename or "upload.bin").strip() or "upload.bin"
    data = await file.read()
    return await _answer_file(
        filename=filename,
        data=data,
        question=question,
        max_tokens=max_tokens,
    )
